"""발표 슬라이드(slides/index.html)를 PowerPoint 파일로 만든다.

사용법: python3 tools/build_pptx.py   (먼저 tools/build_slides.py로 HTML을 만들어 둘 것)
필요: chromium, python-pptx (pip install --user --break-system-packages python-pptx)
결과: slides/발표슬라이드.pptx

- 슬라이드마다 HTML과 똑같은 모양이 나오도록 크롬으로 한 장씩 찍어 그림으로 넣는다
  (코드 글씨 크기 맞추기 등은 슬라이드를 띄울 때 실행되므로 한 장씩 띄워서 찍는다).
- 강사 메모는 PowerPoint 노트로, 슬라이드 제목은 그림의 대체 텍스트로 넣는다.
"""
import html
import re
import subprocess
import tempfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
SLIDES = ROOT / "slides"
OUT = SLIDES / "발표슬라이드.pptx"
W, H = 1600, 900        # 슬라이드 무대 크기 (px)
SCALE = 2               # 2배 해상도로 찍어서 프로젝터에서도 글자가 또렷하게


def text_of(fragment):
    """HTML 조각 → 태그를 뺀 글자"""
    fragment = re.sub(r"<br\s*/?>", " ", fragment)
    return " ".join(html.unescape(re.sub(r"<[^>]+>", "", fragment)).split())


def read_slides(page):
    """(제목, 강사 메모) 목록. 제목(h2)이 없는 표지·구분 슬라이드는 큰 글자(h1)를 제목으로 쓴다"""
    out = []
    for sec in re.findall(r'<section class="slide.*?</section>', page, flags=re.S):
        h2 = re.search(r"<h2>(.*?)</h2>", sec, flags=re.S)
        h1 = re.search(r"<h1>(.*?)</h1>", sec, flags=re.S)
        part = re.search(r'class="divider"><p>(.*?)</p>', sec)   # 구분 슬라이드 위의 작은 글자 (예: 3차시)
        if h2:
            title = text_of(h2.group(1))
        elif h1:
            title = (text_of(part.group(1)) + " · " if part else "") + text_of(h1.group(1))
        else:
            title = ""
        n = re.search(r'<aside class="notes">(.*?)</aside>', sec, flags=re.S)
        notes = text_of(n.group(1)).removeprefix("강사:").strip() if n else ""
        out.append((title, notes))
    return out


def shoot(url, png):
    subprocess.run([
        "chromium", "--headless", "--disable-gpu", "--hide-scrollbars",
        f"--window-size={W},{H}", f"--force-device-scale-factor={SCALE}",
        "--virtual-time-budget=3000", f"--screenshot={png}", url,
    ], check=True, capture_output=True)


def main():
    page = (SLIDES / "index.html").read_text(encoding="utf-8")
    slides = read_slides(page)
    # 캡처용 사본: 아래쪽 번호·단축키 막대를 숨긴다. 그림·라이브러리 경로가 그대로 맞도록 slides 폴더에 잠깐 둔다
    # (<base>로 경로를 바꾸면 슬라이드 스크립트의 주소 바꾸기가 막혀 코드 글씨 맞추기가 안 됨)
    capture = SLIDES / "_capture.html"
    capture.write_text(page.replace("</head>", "<style>#bar { display: none; }</style></head>", 1), encoding="utf-8")
    with tempfile.TemporaryDirectory() as tmp:
        pngs = [Path(tmp) / f"{i:03}.png" for i in range(1, len(slides) + 1)]
        try:
            with ThreadPoolExecutor(3) as pool:   # 세 장씩 동시에 찍는다
                list(pool.map(lambda i: shoot(f"{capture.as_uri()}#{i}", pngs[i - 1]), range(1, len(slides) + 1)))
        finally:
            capture.unlink()

        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9 (33.867cm × 19.05cm)
        blank = prs.slide_layouts[6]
        for (title, notes), png in zip(slides, pngs):
            s = prs.slides.add_slide(blank)
            pic = s.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
            pic.name = title or "슬라이드"
            pic._element.nvPicPr.cNvPr.set("descr", title)   # 대체 텍스트 (화면 읽기 프로그램용)
            if notes:
                s.notes_slide.notes_text_frame.text = notes
        prs.save(OUT)
    print(f"만듦: {OUT.relative_to(ROOT)} (슬라이드 {len(slides)}장, {OUT.stat().st_size / 1e6:.1f}MB)")


if __name__ == "__main__":
    main()
